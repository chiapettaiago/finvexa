from io import BytesIO
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from urllib.parse import parse_qs, urlparse

from werkzeug.security import generate_password_hash

from app import ClaudeAnalysis, Entry, User, UserAccessLog, UserInvitation, claude_entry_data, create_app, db, read_excel


def csrf(client):
    client.get('/login')
    with client.session_transaction() as session:
        return session['csrf']


def test_claude_entry_data_keeps_income_and_expense_types_distinct():
    income = SimpleNamespace(
        description='Salário', kind='receita', category='receita', amount=Decimal('2500.00'),
        paid=True, due=date(2026, 9, 5), expense_date=None,
    )
    fixed_expense = SimpleNamespace(
        description='Aluguel', kind='despesa', category='fixa', amount=Decimal('800.00'),
        paid=False, due=date(2026, 9, 10), expense_date=None,
    )
    daily_expense = SimpleNamespace(
        description='Mercado', kind='despesa', category='dia_a_dia', amount=Decimal('125.40'),
        paid=True, due=None, expense_date=date(2026, 9, 12),
    )

    assert claude_entry_data(income) == {
        'descricao': 'Salário', 'tipo_lancamento': 'receita', 'tipo_despesa': None,
        'valor': '2500.00', 'situacao': 'recebido', 'tipo_data': 'recebimento', 'data': '2026-09-05',
    }
    assert claude_entry_data(fixed_expense)['tipo_despesa'] == 'fixa'
    assert claude_entry_data(fixed_expense)['tipo_data'] == 'vencimento'
    assert claude_entry_data(daily_expense)['tipo_despesa'] == 'dia_a_dia'
    assert claude_entry_data(daily_expense)['tipo_data'] == 'gasto'


def test_read_excel_detects_reordered_columns_and_fills_missing_fields():
    from openpyxl import Workbook

    workbook = BytesIO()
    sheet = Workbook().active
    sheet.append(['Resumo financeiro'])
    sheet.append(['Período: setembro'])
    sheet.append(['Valor', 'Situação', 'Item', 'Competência', 'Natureza'])
    sheet.append([125.50, 'Pendente', 'Internet', 'Setembro', 'Despesa'])
    sheet.append([2500, 'Recebido', 'Salário', '09/2026', 'Entrada'])
    workbook = BytesIO()
    sheet.parent.save(workbook)
    workbook.seek(0)

    entries = read_excel(workbook, 2026)

    assert [(entry['description'], entry['kind'], entry['month'], entry['paid']) for entry in entries] == [
        ('Internet', 'despesa', date(2026, 9, 1), False),
        ('Salário', 'receita', date(2026, 9, 1), True),
    ]
    assert entries[0]['due'] is None


def test_read_excel_detects_monthly_tabs_without_headers():
    from openpyxl import Workbook

    workbook = BytesIO()
    sheet = Workbook().active
    sheet.title = 'Setembro 2024'
    sheet.append(['Salário', 3000, 'Aluguel', 1200, 'PG', None, 'Crédito:', 3000, 'Débito:', 1200])
    sheet.append([None, None, 'Internet', 150, None])
    sheet.append([None, None, None, None, None, None, None, 'Total:', -1200])
    sheet.parent.save(workbook)
    workbook.seek(0)

    entries = read_excel(workbook, 2024)

    assert [(entry['description'], entry['kind'], entry['paid']) for entry in entries] == [
        ('Salário', 'receita', True),
        ('Aluguel', 'despesa', True),
        ('Internet', 'despesa', False),
    ]
    assert all(entry['month'] == date(2024, 9, 1) for entry in entries)


def test_login_create_entry_and_import(tmp_path):
    app = create_app({
        'TESTING': True,
        'SQLALCHEMY_DATABASE_URI': f"sqlite:///{tmp_path / 'test.db'}",
        'SECRET_KEY': 'test-key',
        'RECEIPT_UPLOAD_FOLDER': str(tmp_path / 'receipts'),
    })
    with app.app_context():
        db.create_all()
        db.session.add(User(email='teste@exemplo.com', password=generate_password_hash('senha-segura'), plan='padrao'))
        db.session.commit()
    client = app.test_client()
    favicon = client.get('/favicon.ico')
    assert favicon.status_code == 200
    assert favicon.mimetype == 'image/svg+xml'
    token = csrf(client)
    response = client.post('/login', data={'csrf': token, 'email': 'teste@exemplo.com', 'password': 'senha-segura'}, follow_redirects=True)
    assert response.status_code == 200
    assert b'entry-modal' in response.data
    assert b'import-modal' in response.data
    assert b'mobile-nav' in response.data
    assert b'sidebar-toggle' in response.data
    assert b'id="theme-toggle"' in response.data
    assert b'theme.js' in response.data
    with app.app_context():
        assert db.session.query(UserAccessLog).filter_by(user_id=1).count() == 1
    account_page = client.get('/account')
    assert account_page.status_code == 200
    assert b'Minha conta' in account_page.data
    token = csrf(client)
    response = client.post('/account/profile', data={'csrf': token, 'name': 'Pessoa Teste', 'username': 'pessoa.teste', 'email': 'teste@exemplo.com'}, follow_redirects=True)
    assert b'Pessoa Teste' in response.data
    with app.app_context():
        assert db.session.query(User).one().username == 'pessoa.teste'
        administrator = db.session.query(User).one()
        administrator.is_admin = True
        db.session.commit()
    admin_page = client.get('/admin/users')
    assert admin_page.status_code == 200
    assert b'Criar usu' in admin_page.data
    assert 'Último acesso'.encode() in admin_page.data
    invitation_response = client.post('/admin/invitations', data={'csrf': csrf(client), 'role': 'admin', 'expires_in': '24'}, follow_redirects=True)
    assert invitation_response.status_code == 200
    assert b'Convite criado' in invitation_response.data
    with app.app_context():
        invitation = db.session.query(UserInvitation).one()
        assert invitation.is_admin is True
        invitation_token = invitation.token
    invitation_page = client.get(f'/invite/{invitation_token}')
    assert invitation_page.status_code == 200
    registration = client.post(f'/invite/{invitation_token}', data={'csrf': csrf(client), 'name': 'Pessoa Convidada', 'email': 'convidado@exemplo.com', 'username': 'pessoa.convidada', 'password': 'senha-segura'}, follow_redirects=True)
    assert registration.status_code == 200
    with app.app_context():
        invited_user = db.session.query(User).filter_by(email='convidado@exemplo.com').one()
        assert invited_user.is_admin is True
        assert db.session.get(UserInvitation, invitation.id).used_at is not None
    assert client.get(f'/invite/{invitation_token}').status_code == 410
    client.post('/logout', data={'csrf': csrf(client)})
    client.post('/login', data={'csrf': csrf(client), 'email': 'teste@exemplo.com', 'password': 'senha-segura'}, follow_redirects=True)
    token = csrf(client)
    response = client.post('/admin/users', data={'csrf': token, 'name': 'Novo Usuário', 'username': 'novo.usuario', 'email': 'novo@exemplo.com', 'password': 'senha-segura', 'role': 'admin'}, follow_redirects=True)
    assert response.status_code == 200
    with app.app_context():
        new_user = db.session.query(User).filter_by(username='novo.usuario').one()
        assert new_user is not None
        assert new_user.is_admin is True
        new_user_id = new_user.id
    token = csrf(client)
    response = client.post(f'/admin/users/{new_user_id}/plan', data={'csrf': token, 'plan': 'premium'}, follow_redirects=True)
    assert response.status_code == 200
    with app.app_context():
        assert db.session.get(User, new_user_id).plan == 'premium'
    token = csrf(client)
    client.post(f'/admin/users/{new_user_id}/plan', data={'csrf': token, 'plan': 'invalido'})
    with app.app_context():
        assert db.session.get(User, new_user_id).plan == 'premium'
    token = csrf(client)
    response = client.post('/entry/new', data={'csrf': token, 'description': 'Aluguel', 'kind': 'despesa', 'month': '2026-09', 'due': '2026-09-10', 'amount': '800.00', 'repeat': '1'})
    assert response.status_code == 302
    with app.app_context():
        entry = db.session.query(Entry).one()
        entry_id = entry.id
    token = csrf(client)
    response = client.post('/entry/new', data={'csrf': token, 'description': 'Internet', 'kind': 'despesa', 'category': 'fixa', 'month': '2026-09', 'due': '2026-09-15', 'amount': '100.00', 'repeat': 'never'})
    assert response.status_code == 302
    with app.app_context():
        assert db.session.query(Entry).filter_by(description='Internet').count() == 1
    page = client.get('/?month=2026-09')
    assert b'<option value="never">Nunca</option>' in page.data
    token = csrf(client)
    receipt = BytesIO(b'%PDF-1.4\ncomprovante de teste')
    response = client.post(f'/entry/{entry_id}/edit', data={'csrf': token, 'description': 'Aluguel', 'kind': 'despesa', 'category': 'fixa', 'month': '2026-09', 'due': '2026-09-10', 'amount': '800.00', 'receipt': (receipt, 'aluguel.pdf')}, content_type='multipart/form-data')
    assert response.status_code == 302
    downloaded = client.get(f'/entry/{entry_id}/receipt')
    assert downloaded.status_code == 200
    assert downloaded.mimetype == 'application/pdf'
    assert downloaded.data.startswith(b'%PDF-')
    assert downloaded.headers['X-Frame-Options'] == 'SAMEORIGIN'
    attachment = client.get(f'/entry/{entry_id}/receipt?download=1')
    assert 'attachment' in attachment.headers['Content-Disposition']
    token = csrf(client)
    large_receipt = BytesIO(b'\xff\xd8\xff' + b'x' * (7 * 1024 * 1024))
    response = client.post(f'/entry/{entry_id}/edit', data={'csrf': token, 'description': 'Aluguel', 'kind': 'despesa', 'category': 'fixa', 'month': '2026-09', 'due': '2026-09-10', 'amount': '800.00', 'receipt': (large_receipt, 'foto.jpg')}, content_type='multipart/form-data')
    assert response.status_code == 302
    token = csrf(client)
    oversized_receipt = BytesIO(b'\xff\xd8\xff' + b'x' * (22 * 1024 * 1024))
    response = client.post(f'/entry/{entry_id}/edit', data={'csrf': token, 'description': 'Aluguel', 'kind': 'despesa', 'category': 'fixa', 'month': '2026-09', 'due': '2026-09-10', 'amount': '800.00', 'receipt': (oversized_receipt, 'foto-grande.jpg')}, content_type='multipart/form-data')
    assert response.status_code == 413
    assert b'Arquivo muito grande' in response.data
    page = client.get('/?month=2026-09')
    assert b'Pagar' in page.data
    reports = client.get('/reports?month=2026-09')
    assert b'Pagar' in reports.data
    assert b'Excel' in reports.data
    assert b'PDF' in reports.data
    excel_report = client.get('/reports/export/excel?month=2026-09')
    assert excel_report.status_code == 200
    assert excel_report.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert excel_report.data.startswith(b'PK')
    pdf_report = client.get('/reports/export/pdf?month=2026-09')
    assert pdf_report.status_code == 200
    assert pdf_report.mimetype == 'application/pdf'
    assert pdf_report.data.startswith(b'%PDF-')
    token = csrf(client)
    shared = client.post('/reports/share', data={'csrf': token, 'month': '2026-09'})
    assert shared.status_code == 302
    share_token = parse_qs(urlparse(shared.headers['Location']).query)['share'][0]
    assert '2026-09' not in share_token
    shared_page = client.get(f'/shared-report/{share_token}')
    assert shared_page.status_code == 200
    assert b'Relat' in shared_page.data
    assert b'Aluguel' in shared_page.data
    assert b'Pessoa Teste' in shared_page.data
    assert b'Saldo realizado' in shared_page.data
    assert b'id="theme-toggle"' not in shared_page.data
    assert b'theme.js' not in shared_page.data
    links_page = client.get('/shared-links')
    assert links_page.status_code == 200
    assert b'Links compartilh' in links_page.data
    assert b'2h' in links_page.data
    assert f'data-confirm-action="/entry/{entry_id}/toggle"'.encode() in reports.data
    assert reports.data.count(b'id="confirm-modal"') == 1
    assert b'id="receipt-open"' in reports.data
    assert b'Ver comprovante' in page.data
    token = csrf(client)
    client.post(f'/entry/{entry_id}/toggle', data={'csrf': token})
    page = client.get('/?month=2026-09')
    assert b'Pago' in page.data
    shared_page = client.get(f'/shared-report/{share_token}')
    assert b'Ver comprovante' in shared_page.data
    assert b'id="receipt-modal"' in shared_page.data
    shared_receipt = client.get(f'/shared-report/{share_token}/receipt/{entry_id}')
    assert shared_receipt.status_code == 200
    assert shared_receipt.mimetype == 'image/jpeg'
    with app.app_context():
        db.session.add(ClaudeAnalysis(user_id=1, month=entry.month, content='Priorize o aluguel.', model='modelo-teste'))
        db.session.commit()
    page = client.get('/?month=2026-09')
    assert b'Priorize o aluguel.' in page.data
    assert b'Hist\xc3\xb3rico de an\xc3\xa1lises' in page.data
    assert b'Assistente Claude' not in page.data
    assert b'Modelo:' not in page.data
    token = csrf(client)
    response = client.post('/api/suggestions', data={'csrf': token, 'month': '2026-09'})
    assert response.status_code in (200, 503)
    token = csrf(client)
    client.post('/entry/new', data={'csrf': token, 'description': 'Academia', 'kind': 'despesa', 'month': '2026-09', 'due': '2026-09-08', 'amount': '90.00', 'repeat': '1'})
    page = client.get('/?month=2026-09')
    assert page.data.index(b'Academia') < page.data.index(b'Aluguel')
    assert b'Receitas' in page.data
    assert b'Despesas fixas' in page.data
    assert b'Despesas do dia a dia' in page.data
    reports = client.get('/reports?month=2026-09')
    assert reports.status_code == 200
    assert b'Relat\xc3\xb3rios financeiros' in reports.data
    assert b'annual-chart' in reports.data
    annual_planning = client.get('/planejamento-anual?month=2026-09')
    assert annual_planning.status_code == 200
    assert b'Planejamento anual' in annual_planning.data
    assert b'Academia' in annual_planning.data
    token = csrf(client)
    client.post('/entry/new', data={'csrf': token, 'description': 'Mercado', 'kind': 'despesa', 'category': 'dia_a_dia', 'month': '2026-08', 'entry_date': '2026-09-12', 'amount': '125.40', 'repeat': '12'})
    with app.app_context():
        daily = db.session.query(Entry).filter_by(description='Mercado').one()
        assert daily.category == 'dia_a_dia'
        assert daily.month.isoformat() == '2026-09-01'
        assert daily.expense_date.isoformat() == '2026-09-12'
        assert daily.due is None
    token = csrf(client)
    client.post('/entry/new', data={'csrf': token, 'description': 'Uber', 'kind': 'despesa', 'category': 'dia_a_dia', 'month': '2026-09', 'entry_date': '2026-09-05', 'amount': '13.27', 'repeat': 'never'})
    page = client.get('/?month=2026-09')
    assert page.data.index(b'Uber') < page.data.index(b'Mercado')
    reports = client.get('/reports?month=2026-09')
    assert reports.data.index(b'Uber') < reports.data.index(b'Mercado')
    token = csrf(client)
    workbook = BytesIO()
    from openpyxl import Workbook
    ws = Workbook().active
    ws.append(['Descrição', 'Tipo', 'Mês', 'Vencimento', 'Valor', 'Pago'])
    ws.append(['Salário', 'receita', '2026-09', '05/09/2026', 2500, 'Sim'])
    ws.parent.save(workbook)
    workbook.seek(0)
    response = client.post('/import', data={'csrf': token, 'year': '2026', 'file': (workbook, 'entrada.xlsx')}, content_type='multipart/form-data')
    assert response.status_code == 302
    invalid_workbook = BytesIO()
    invalid_sheet = Workbook().active
    invalid_sheet.append(['Coluna inválida'])
    invalid_sheet.parent.save(invalid_workbook)
    invalid_workbook.seek(0)
    response = client.post(
        '/import',
        data={'csrf': csrf(client), 'year': '2026', 'file': (invalid_workbook, 'invalida.xlsx')},
        content_type='multipart/form-data',
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert 'Formato não reconhecido.' in response.text
    assert 'O que alterar no Excel:' in response.text
    assert 'Descrição, Tipo, Mês, Vencimento, Valor e Pago.' in response.text
    assert b'flash error import-error' in response.data
    with app.app_context():
        assert db.session.query(Entry).count() == 6
        deleted_entry = Entry(user_id=1, description='Apagar depois', kind='despesa', category='fixa', month=date(2026, 9, 1), amount=Decimal('10.00'))
        db.session.add(deleted_entry)
        db.session.commit()
        deleted_entry_id = deleted_entry.id
    response = client.post(f'/entry/{deleted_entry_id}/delete', data={'csrf': csrf(client)}, follow_redirects=True)
    assert response.status_code == 200
    assert 'Lançamento excluído.'.encode() in response.data
    assert b'Apagar depois' not in response.data
    with app.app_context():
        deleted_entry = db.session.get(Entry, deleted_entry_id)
        assert deleted_entry is not None
        assert deleted_entry.deleted_at is not None
    assert client.post(f'/entry/{deleted_entry_id}/toggle', data={'csrf': csrf(client)}).status_code == 404
    assert client.get(f'/entry/{deleted_entry_id}/receipt').status_code == 404


def test_plan_limits_share_links(tmp_path):
    app = create_app({'TESTING': True, 'SQLALCHEMY_DATABASE_URI': f'sqlite:///{tmp_path / "planos.db"}', 'RECEIPT_UPLOAD_FOLDER': str(tmp_path / 'receipts')})
    with app.app_context():
        db.create_all()
        db.session.add(User(email='p@exemplo.com', password=generate_password_hash('senha-segura'), plan='basico'))
        db.session.commit()
    client = app.test_client()
    client.post('/login', data={'csrf': csrf(client), 'email': 'p@exemplo.com', 'password': 'senha-segura'})
    blocked = client.post('/reports/share', data={'csrf': csrf(client), 'month': '2026-09'}, follow_redirects=True)
    assert 'Básico não inclui'.encode() in blocked.data
    with app.app_context():
        user = db.session.query(User).one()
        assert user.plan_info['ai'] == 2
        user.plan = 'padrao'
        db.session.commit()
    for _ in range(5):
        assert 'share=' in client.post('/reports/share', data={'csrf': csrf(client), 'month': '2026-09'}).headers['Location']
    over = client.post('/reports/share', data={'csrf': csrf(client), 'month': '2026-09'}, follow_redirects=True)
    assert 'limite de 5 links'.encode() in over.data


def test_landing_page_for_visitors_and_dashboard_for_users(tmp_path):
    app = create_app({'TESTING': True, 'SQLALCHEMY_DATABASE_URI': f'sqlite:///{tmp_path / "landing.db"}', 'RECEIPT_UPLOAD_FOLDER': str(tmp_path / 'receipts')})
    with app.app_context():
        db.create_all()
        db.session.add(User(email='l@exemplo.com', password=generate_password_hash('senha-segura')))
        db.session.commit()
    client = app.test_client()
    landing = client.get('/')
    assert landing.status_code == 200
    assert b'Sou cliente' in landing.data and b'href="/login"' in landing.data
    assert b'49,90' in landing.data and b'style="' not in landing.data
    client.post('/login', data={'csrf': csrf(client), 'email': 'l@exemplo.com', 'password': 'senha-segura'})
    assert b'Sou cliente' not in client.get('/').data


def test_subscription_creates_account_only_after_payment(tmp_path, monkeypatch):
    import app as app_module
    from app import Subscription
    monkeypatch.setenv('MP_ACCESS_TOKEN', 'token-teste')
    monkeypatch.delenv('MP_WEBHOOK_SECRET', raising=False)
    remote = {}

    def fake_mp(method, path, payload=None):
        if method == 'POST':
            remote.update(payload, id='pre-1', status='pending')
            return {'id': 'pre-1', 'init_point': 'https://www.mercadopago.com.br/checkout/pre-1'}
        return dict(remote)

    monkeypatch.setattr(app_module, 'mp_request', fake_mp)
    app = create_app({'TESTING': True, 'SQLALCHEMY_DATABASE_URI': f'sqlite:///{tmp_path / "assinatura.db"}', 'RECEIPT_UPLOAD_FOLDER': str(tmp_path / 'receipts')})
    with app.app_context():
        db.create_all()
    client = app.test_client()
    assert client.get('/assinar/inexistente').status_code == 404
    page = client.post('/assinar/padrao', data={'csrf': csrf(client), 'email': 'Cliente@Exemplo.com'})
    assert page.status_code == 200 and b'checkout/pre-1' in page.data
    assert remote['auto_recurring']['transaction_amount'] == 49.9
    with app.app_context():
        subscription = db.session.query(Subscription).one()
        token = subscription.token
        assert subscription.status == 'pending' and subscription.email == 'cliente@exemplo.com'
    # Pagamento pendente: não há cadastro e nenhuma conta é criada.
    assert client.get('/cadastro/' + token).status_code == 410
    assert b'Aguardando' in client.get('/assinatura/retorno?ref=' + token).data
    with app.app_context():
        assert db.session.query(User).count() == 0
    # Pagamento confirmado no Mercado Pago libera o cadastro com o plano contratado.
    remote['status'] = 'authorized'
    back = client.get('/assinatura/retorno?ref=' + token)
    assert back.status_code == 302 and back.headers['Location'].endswith('/cadastro/' + token)
    done = client.post('/cadastro/' + token, data={'csrf': csrf(client), 'name': 'Cliente', 'username': 'cliente', 'password': 'senha-bem-longa'})
    assert done.status_code == 302
    with app.app_context():
        user = db.session.query(User).one()
        assert user.plan == 'padrao' and user.email == 'cliente@exemplo.com'
    assert client.get('/cadastro/' + token).status_code == 410
    # Cancelamento via webhook bloqueia o login.
    remote['status'] = 'cancelled'
    assert client.post('/webhooks/mercadopago', json={'data': {'id': 'pre-1'}}).status_code == 200
    client.post('/logout', data={'csrf': csrf(client)})
    blocked = client.post('/login', data={'csrf': csrf(client), 'email': 'cliente', 'password': 'senha-bem-longa'})
    assert 'assinatura está inativa'.encode() in blocked.data


def test_subscription_unavailable_without_token_and_webhook_signature(tmp_path, monkeypatch):
    monkeypatch.delenv('MP_ACCESS_TOKEN', raising=False)
    monkeypatch.setenv('MP_WEBHOOK_SECRET', 'segredo')
    app = create_app({'TESTING': True, 'SQLALCHEMY_DATABASE_URI': f'sqlite:///{tmp_path / "sem.db"}', 'RECEIPT_UPLOAD_FOLDER': str(tmp_path / 'receipts')})
    with app.app_context():
        db.create_all()
    client = app.test_client()
    assert client.get('/assinar/basico').status_code == 503
    assert client.post('/webhooks/mercadopago', json={'data': {'id': 'x'}}, headers={'x-signature': 'ts=1,v1=errado'}).status_code == 401
