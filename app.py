import calendar
import base64
import hashlib
import hmac
import json
import os
import re
import secrets
import unicodedata
import urllib.error
import urllib.request
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from functools import wraps
from io import BytesIO

import click
from dotenv import load_dotenv
from sqlalchemy.engine import URL
from cryptography.fernet import Fernet, InvalidToken
from flask import Flask, abort, flash, g, jsonify, redirect, render_template, request, send_file, send_from_directory, session, url_for
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import func, inspect, or_, select, text
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from werkzeug.middleware.proxy_fix import ProxyFix
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

load_dotenv()
db = SQLAlchemy()
MONTHS = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]


PLANS = {
    "basico": {"name": "Básico", "price": Decimal("19.90"), "ai": 2, "share_links": 0},
    "padrao": {"name": "Padrão", "price": Decimal("49.90"), "ai": 10, "share_links": 5},
    "premium": {"name": "Premium", "price": Decimal("79.90"), "ai": 15, "share_links": None},
}
DEFAULT_PLAN = "basico"


def format_brl(value):
    return "R$ " + f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def entry_report_row(entry):
    daily = entry.kind == "despesa" and entry.category == "dia_a_dia"
    entry_date = entry.expense_date if daily else entry.due
    return (
        entry.description,
        "Receita" if entry.kind == "receita" else "Despesa",
        "—" if entry.kind == "receita" else "Dia a dia" if daily else "Fixa",
        entry_date.strftime("%d/%m/%Y") if entry_date else "Sem data",
        entry.amount,
        "Recebido" if entry.kind == "receita" and entry.paid else "Pago" if entry.paid else "Pendente",
    )

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(254), unique=True, nullable=False)
    name = db.Column(db.String(120))
    username = db.Column(db.String(50), unique=True)
    password = db.Column(db.String(255), nullable=False)
    avatar_file = db.Column(db.String(80))
    is_admin = db.Column(db.Boolean, nullable=False, default=False)
    ai_monthly_limit = db.Column(db.Integer, nullable=False, default=10)
    plan = db.Column(db.String(20), nullable=False, default=DEFAULT_PLAN)

    @property
    def plan_info(self):
        return PLANS.get(self.plan) or PLANS[DEFAULT_PLAN]

class Entry(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    description = db.Column(db.String(150), nullable=False)
    kind = db.Column(db.String(10), nullable=False)
    category = db.Column(db.String(20), nullable=False, default="fixa")
    month = db.Column(db.Date, nullable=False, index=True)
    due = db.Column(db.Date)
    expense_date = db.Column(db.Date)
    amount = db.Column(db.Numeric(12, 2), nullable=False)
    paid = db.Column(db.Boolean, nullable=False, default=False)
    receipt_file = db.Column(db.String(80))
    receipt_name = db.Column(db.String(255))
    receipt_mimetype = db.Column(db.String(100))
    receipt_uploaded_at = db.Column(db.DateTime(timezone=True))
    deleted_at = db.Column(db.DateTime(timezone=True), index=True)

class ClaudeAnalysis(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    month = db.Column(db.Date, nullable=False, index=True)
    content = db.Column(db.Text, nullable=False)
    model = db.Column(db.String(100), nullable=False)
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=lambda: datetime.now(timezone.utc), index=True)


class SharedReport(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    month = db.Column(db.Date, nullable=False, index=True)
    token = db.Column(db.String(500), nullable=False, unique=True)
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=lambda: datetime.now(timezone.utc), index=True)
    expires_at = db.Column(db.DateTime(timezone=True), nullable=False, index=True)


class UserInvitation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(254), nullable=False, index=True)
    is_admin = db.Column(db.Boolean, nullable=False, default=False)
    token = db.Column(db.String(255), nullable=False, unique=True)
    invited_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=lambda: datetime.now(timezone.utc))
    expires_at = db.Column(db.DateTime(timezone=True), nullable=False, index=True)
    used_at = db.Column(db.DateTime(timezone=True), index=True)


class UserAccessLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    accessed_at = db.Column(db.DateTime(timezone=True), nullable=False, default=lambda: datetime.now(timezone.utc), index=True)


class Subscription(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    token = db.Column(db.String(64), nullable=False, unique=True)
    plan = db.Column(db.String(20), nullable=False)
    email = db.Column(db.String(254), nullable=False, index=True)
    mp_preapproval_id = db.Column(db.String(64), index=True)
    status = db.Column(db.String(20), nullable=False, default="pending")
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), index=True)
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=lambda: datetime.now(timezone.utc))


def mp_request(method, path, payload=None):
    """Chamada à API do Mercado Pago; isolada para poder ser substituída nos testes."""
    access_token = os.environ.get("MP_ACCESS_TOKEN")
    if not access_token:
        raise RuntimeError("Mercado Pago não configurado.")
    request_ = urllib.request.Request(
        "https://api.mercadopago.com" + path,
        data=json.dumps(payload).encode() if payload is not None else None,
        method=method,
        headers={"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"},
    )
    try:
        with urllib.request.urlopen(request_, timeout=15) as response:
            return json.loads(response.read())
    except urllib.error.HTTPError as error:
        raise RuntimeError(f"Mercado Pago respondeu {error.code}: {error.read().decode(errors='replace')[:500]}") from error


def utc_datetime(value):
    return value.replace(tzinfo=timezone.utc) if value and value.tzinfo is None else value

def money(value):
    if value is None or isinstance(value, bool):
        raise ValueError("Informe um valor monetário válido.")
    s = str(value).replace("R$", "").replace(" ", "").strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        result = Decimal(s).quantize(Decimal("0.01"))
        if not result.is_finite() or result < 0 or result >= Decimal("10000000000"):
            raise ValueError()
        return result
    except (InvalidOperation, ValueError):
        raise ValueError("Valor inválido: use um número positivo com até 10 dígitos inteiros.")

def parse_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Data inválida: {value}. Use DD/MM/AAAA.")

def parse_paid(value):
    s = str(value or "").strip().lower()
    if s not in ("", "sim", "não", "nao", "true", "false", "1", "0", "pago", "quitado", "recebido", "pg", "pendente", "aberto"):
        raise ValueError("Pagamento deve ser Sim ou Não.")
    return s in ("sim", "true", "1", "pago", "quitado", "recebido", "pg")


def claude_entry_data(entry):
    """Converte um lançamento para o contrato de contexto enviado ao Claude."""
    is_expense = entry.kind == "despesa"
    is_daily_expense = is_expense and entry.category == "dia_a_dia"
    event_date = entry.expense_date if is_daily_expense else entry.due
    return {
        "descricao": entry.description,
        "tipo_lancamento": entry.kind,
        # Receitas não têm tipo de despesa. Não reutilize a categoria interna
        # "receita" para esse campo, pois isso induz o modelo a classificá-la
        # como uma despesa.
        "tipo_despesa": entry.category if is_expense else None,
        "valor": str(entry.amount),
        "situacao": "pago" if is_expense and entry.paid else "recebido" if entry.paid else "pendente",
        "tipo_data": "gasto" if is_daily_expense else "vencimento" if is_expense else "recebimento",
        "data": event_date.isoformat() if event_date else None,
    }

RECEIPT_TYPES = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}

def save_receipt(upload, folder):
    original = secure_filename(upload.filename or "")
    extension = os.path.splitext(original)[1].lower()
    if not original or extension not in RECEIPT_TYPES:
        raise ValueError("Envie um comprovante em PDF, PNG, JPG ou WEBP.")
    data = upload.read()
    if not data or len(data) > 20 * 1024 * 1024:
        raise ValueError("O comprovante deve ter entre 1 byte e 20 MB.")
    signatures = {
        ".pdf": data.startswith(b"%PDF-"),
        ".png": data.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": data.startswith(b"\xff\xd8\xff"),
        ".jpeg": data.startswith(b"\xff\xd8\xff"),
        ".webp": data.startswith(b"RIFF") and data[8:12] == b"WEBP",
    }
    if not signatures[extension]:
        raise ValueError("O conteúdo do comprovante não corresponde ao tipo do arquivo.")
    stored = f"{uuid.uuid4().hex}{extension}"
    os.makedirs(folder, mode=0o700, exist_ok=True)
    with open(os.path.join(folder, stored), "wb") as destination:
        destination.write(data)
    return stored, original, RECEIPT_TYPES[extension]

AVATAR_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}

def save_avatar(upload, folder):
    original = secure_filename(upload.filename or "")
    extension = os.path.splitext(original)[1].lower()
    if not original or extension not in AVATAR_TYPES:
        raise ValueError("Envie uma foto em PNG, JPG ou WEBP.")
    data = upload.read()
    if not data or len(data) > 20 * 1024 * 1024:
        raise ValueError("A foto deve ter entre 1 byte e 20 MB.")
    signatures = {
        ".png": data.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": data.startswith(b"\xff\xd8\xff"),
        ".jpeg": data.startswith(b"\xff\xd8\xff"),
        ".webp": data.startswith(b"RIFF") and data[8:12] == b"WEBP",
    }
    if not signatures[extension]:
        raise ValueError("O conteúdo da foto não corresponde ao tipo do arquivo.")
    stored = f"{uuid.uuid4().hex}{extension}"
    os.makedirs(folder, mode=0o700, exist_ok=True)
    with open(os.path.join(folder, stored), "wb") as destination:
        destination.write(data)
    return stored, AVATAR_TYPES[extension]


def excel_label(value):
    """Normaliza cabeçalhos para aceitar nomes e acentuação diferentes."""
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(char for char in text if not unicodedata.combining(char)).lower().split())


def excel_month(value, fallback_year):
    if not value:
        return date(fallback_year, date.today().month, 1)
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    text = excel_label(value)
    for month_number, month_name in enumerate(MONTHS, 1):
        if text in (excel_label(month_name), str(month_number), f"{month_number:02d}"):
            return date(fallback_year, month_number, 1)
    for fmt in ("%Y-%m", "%m/%Y", "%Y/%m", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            parsed = datetime.strptime(str(value).strip(), fmt)
            return date(parsed.year, parsed.month, 1)
        except ValueError:
            pass
    raise ValueError("Mês inválido: use AAAA-MM, MM/AAAA ou o nome do mês.")


def find_excel_columns(rows):
    aliases = {
        "description": {"descricao", "item", "lancamento", "nome", "conta"},
        "kind": {"tipo", "natureza", "tipo de lancamento", "categoria"},
        "month": {"mes", "mes de referencia", "competencia", "referencia", "periodo"},
        "due": {"vencimento", "data", "data de vencimento", "data pagamento", "data de pagamento"},
        "amount": {"valor", "valor r", "valor rs", "preco", "quantia", "total"},
        "paid": {"pago", "paga", "pagamento", "situacao", "status", "quitado"},
    }
    for row_number, row in enumerate(rows[:30], 1):
        columns = {}
        for index, value in enumerate(row):
            label = excel_label(value)
            for field, names in aliases.items():
                if label in names and field not in columns:
                    columns[field] = index
        if {"description", "amount"}.issubset(columns):
            return row_number, columns
    return None, None


def sheet_month(title):
    match = re.search(r"(.*?)(?:\s+)(\d{4})$", str(title).strip())
    if not match:
        return None
    name, year = excel_label(match.group(1)), int(match.group(2))
    for month_number, month_name in enumerate(MONTHS, 1):
        normalized_month = excel_label(month_name)
        if name == normalized_month or name[:3] == normalized_month[:3]:
            return date(year, month_number, 1)
    return None


def read_monthly_tabs(workbook):
    """Lê o layout de abas mensais, com pares descrição/valor sem cabeçalhos."""
    entries = []
    ignored_labels = {"credito", "debito", "total", "saldo"}
    for ws in workbook.worksheets:
        month = sheet_month(ws.title)
        if not month:
            continue
        for row in ws.iter_rows(values_only=True):
            for column in range(0, len(row) - 1, 2):
                description, amount = row[column], row[column + 1]
                label = excel_label(description).rstrip(":")
                if not label or label in ignored_labels or amount in (None, ""):
                    continue
                try:
                    value = money(amount)
                except ValueError:
                    continue
                kind = "receita" if column == 0 else "despesa"
                paid = True if kind == "receita" else parse_paid(row[column + 2] if column + 2 < len(row) else None)
                entries.append(dict(description=str(description).strip(), kind=kind, month=month, due=None, amount=value, paid=paid))
    return entries


def read_excel(file, year):
    wb = load_workbook(file, data_only=True, read_only=True)
    try:
        monthly_entries = read_monthly_tabs(wb)
        if monthly_entries:
            return monthly_entries
        ws = wb.active
        if ws.max_row > 5000 or ws.max_column > 100:
            raise ValueError("Planilha muito grande: limite de 5.000 linhas e 100 colunas.")
        rows = list(ws.values)
        if not rows:
            raise ValueError("Planilha vazia.")
        result = []
        def add(description, kind, month, due, amount, paid):
            description = str(description or "").strip()
            if not description or len(description) > 150:
                raise ValueError("Descrição obrigatória, com até 150 caracteres.")
            result.append(dict(description=description, kind=kind, month=month, due=parse_date(due), amount=money(amount), paid=parse_paid(paid)))
        header_row, columns = find_excel_columns(rows)
        if columns:
            for number, row in enumerate(rows[header_row:], header_row + 1):
                if not any(v is not None for v in row):
                    continue
                try:
                    value = lambda field: row[columns[field]] if columns.get(field, len(row)) < len(row) else None
                    desc, amount = value("description"), value("amount")
                    if not str(desc or "").strip() or amount in (None, ""):
                        continue
                    kind = excel_label(value("kind"))
                    kind = "receita" if kind in ("receita", "entrada", "recebimento", "ganho") else "despesa"
                    month = excel_month(value("month"), year)
                    due, paid = value("due"), value("paid")
                    add(desc, kind, month, due, amount, paid)
                except (ValueError, TypeError) as exc:
                    raise ValueError(f"Linha {number}: {exc}")
        else:
            groups = []
            for row in rows[:3]:
                for col, value in enumerate(row):
                    name = str(value or "").strip().lower()
                    if name in [m.lower() for m in MONTHS] and col+2 < len(row):
                        # The merged month title starts at the due-date column.
                        if any("vencimento" in str(v).lower() for r in rows[:3] for v in r):
                            if not any(g[1] == MONTHS.index(name.capitalize())+1 for g in groups):
                                groups.append((col, MONTHS.index(name.capitalize())+1))
                if groups:
                    break
            if not groups:
                raise ValueError("Formato não reconhecido. Baixe o modelo ou use o layout da imagem.")
            for number, row in enumerate(rows[2:], 3):
                desc = str(row[0] or "").strip() if row else ""
                if not desc or desc.lower().rstrip(":") in ("total", "saldo"):
                    continue
                for col, month in groups:
                    if col+2 >= len(row) or row[col+1] is None:
                        continue
                    try:
                        income = desc.lower().rstrip(":") == "recebimentos"
                        add(desc, "receita" if income else "despesa", date(year, month, 1), None if income else row[col], row[col+1], "Sim" if income else row[col+2])
                    except ValueError as exc:
                        raise ValueError(f"Linha {number}, {MONTHS[month-1]}: {exc}")
        if not result:
            raise ValueError("Nenhum lançamento encontrado. Fórmulas precisam estar calculadas e salvas no Excel.")
        return result
    finally:
        wb.close()


def import_error_message(error):
    """Traduz erros de importação em uma orientação acionável para a planilha."""
    detail = str(error)
    guidance = "Revise a célula indicada, salve a planilha e tente importar novamente."

    if "Descrição obrigatória" in detail:
        guidance = "Preencha a coluna Descrição com até 150 caracteres."
    elif "Tipo deve ser" in detail:
        guidance = "Na coluna Tipo, informe apenas 'despesa' ou 'receita'."
    elif "Data inválida" in detail:
        guidance = "Na coluna Vencimento, use uma data no formato DD/MM/AAAA."
    elif "Valor inválido" in detail or "valor monetário válido" in detail:
        guidance = "Na coluna Valor, informe um número positivo, por exemplo 125,50."
    elif "Pagamento deve ser" in detail:
        guidance = "Na coluna Pago, use apenas Sim ou Não."
    elif "Formato não reconhecido" in detail:
        guidance = "Use o modelo disponível no link 'Baixar modelo' ou mantenha as colunas: Descrição, Tipo, Mês, Vencimento, Valor e Pago."
    elif "Planilha vazia" in detail:
        guidance = "Inclua o cabeçalho e pelo menos um lançamento antes de salvar o arquivo."
    elif "Nenhum lançamento encontrado" in detail:
        guidance = "Preencha pelo menos um lançamento com valores, salve o arquivo e tente novamente."
    elif "Planilha muito grande" in detail:
        guidance = "Reduza a planilha para, no máximo, 5.000 linhas e 100 colunas."
    elif "Selecione um arquivo" in detail:
        guidance = "Envie um arquivo com a extensão .xlsx."

    return f"{detail} O que alterar no Excel: {guidance}"


def create_app(config=None):
    app = Flask(__name__)
    os.makedirs(app.instance_path, exist_ok=True)
    key_path = os.path.join(app.instance_path, "secret.key")
    if not os.environ.get("SECRET_KEY") and not os.path.exists(key_path):
        with open(key_path, "w", opener=lambda p, f: os.open(p, f, 0o600)) as f:
            f.write(secrets.token_hex(32))
    database_url = os.environ.get("DATABASE_URL")
    if not database_url and os.environ.get("DB_HOST"):
        database_url = URL.create(
            "mysql+pymysql",
            username=os.environ.get("DB_USER"),
            password=os.environ.get("DB_PASS"),
            host=os.environ["DB_HOST"].strip(),
            port=int(os.environ.get("DB_PORT", "3306")),
            database=os.environ.get("DB_NAME"),
            query={"charset": os.environ.get("DB_CHARSET", "utf8mb4")},
        ).render_as_string(hide_password=False)
    database_url = database_url or "sqlite:///financas.db"
    engine_options = {"pool_pre_ping": True}
    if database_url.startswith("mysql"):
        engine_options.update(pool_size=3, max_overflow=2, pool_recycle=1800, connect_args={"connect_timeout": int(os.environ.get("DB_CONNECT_TIMEOUT", "5"))})
    app.config.update(SECRET_KEY=os.environ.get("SECRET_KEY") or open(key_path).read(), APP_NAME=os.environ.get("APP_NAME", "Finvexa"), APP_SLOGAN=os.environ.get("APP_SLOGAN", "Seus números, decisões mais inteligentes."), SQLALCHEMY_DATABASE_URI=database_url, SQLALCHEMY_ENGINE_OPTIONS=engine_options, SQLALCHEMY_TRACK_MODIFICATIONS=False, MAX_CONTENT_LENGTH=21*1024*1024, SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax", SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE") == "true")
    app.config.setdefault("RECEIPT_UPLOAD_FOLDER", os.path.join(app.instance_path, "receipts"))
    app.config.setdefault("AVATAR_UPLOAD_FOLDER", os.path.join(app.instance_path, "avatars"))
    if config:
        app.config.update(config)
        if "SQLALCHEMY_DATABASE_URI" in config and "SQLALCHEMY_ENGINE_OPTIONS" not in config and not config["SQLALCHEMY_DATABASE_URI"].startswith("mysql"):
            app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"pool_pre_ping": True}
    share_key = base64.urlsafe_b64encode(hashlib.sha256(app.config["SECRET_KEY"].encode()).digest())
    app.config["REPORT_SHARE_CIPHER"] = Fernet(share_key)
    if not app.testing:
        from logging.handlers import RotatingFileHandler
        file_handler = RotatingFileHandler(os.path.join(app.instance_path, "app.log"), maxBytes=1_000_000, backupCount=3, encoding="utf-8")
        file_handler.setLevel("WARNING")
        file_handler.setFormatter(__import__("logging").Formatter("%(asctime)s %(levelname)s %(message)s"))
        app.logger.addHandler(file_handler)
    db.init_app(app)
    if os.environ.get("TRUST_PROXY") == "true":
        app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_prefix=1)
    app.jinja_env.filters["brl"] = format_brl
    def br_datetime(value):
        if not value:
            return ""
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        from zoneinfo import ZoneInfo
        return value.astimezone(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y às %H:%M")
    app.jinja_env.filters["br_datetime"] = br_datetime
    @app.errorhandler(413)
    def upload_too_large(error):
        return render_template("upload_error.html"), 413

    @app.before_request
    def load_current_user():
        g.current_user = db.session.get(User, session["user_id"]) if session.get("user_id") else None
    @app.context_processor
    def context():
        session.setdefault("csrf", secrets.token_hex(32))
        return dict(csrf=session["csrf"], months=MONTHS, today=date.today(), current_user=g.current_user, plans=PLANS, app_name=app.config["APP_NAME"], app_slogan=app.config["APP_SLOGAN"])
    @app.before_request
    def csrf_check():
        if request.method == "POST" and request.endpoint != "mercadopago_webhook" and not secrets.compare_digest(session.get("csrf", "missing"), request.form.get("csrf", "")):
            abort(400, "Sessão expirada. Atualize a página e tente novamente.")
    @app.after_request
    def headers(response):
        response.headers["X-Content-Type-Options"] = "nosniff"
        receipt_preview = request.endpoint in ("receipt", "shared_receipt") and request.args.get("download") != "1"
        response.headers["X-Frame-Options"] = "SAMEORIGIN" if receipt_preview else "DENY"
        policy = "default-src self; style-src self; form-action self; frame-ancestors self" if receipt_preview else "default-src self; style-src self; form-action self; frame-ancestors none"
        response.headers["Content-Security-Policy"] = policy.replace("self", "\u0027self\u0027").replace("none", "\u0027none\u0027")
        response.headers["Cache-Control"] = "no-store"
        return response
    @app.get("/healthz")
    def healthz():
        return {"status": "ok"}
    @app.get("/favicon.ico")
    def favicon():
        return send_from_directory(app.static_folder, "favicon.svg", mimetype="image/svg+xml")
    def login_required(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            if not session.get("user_id"):
                return redirect(url_for("login"))
            return fn(*args, **kwargs)
        return wrapper
    def admin_required(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            if not session.get("user_id"):
                return redirect(url_for("login"))
            if not g.current_user or not g.current_user.is_admin:
                abort(403)
            return fn(*args, **kwargs)
        return wrapper
    def sync_subscription(subscription):
        """Consulta o Mercado Pago e grava a situação real da assinatura; nunca confia em dados vindos do navegador."""
        data = mp_request("GET", f"/preapproval/{subscription.mp_preapproval_id}")
        if data.get("external_reference") != subscription.token:
            raise ValueError("Assinatura não corresponde ao pedido.")
        if data.get("status") in ("pending", "authorized", "paused", "cancelled"):
            subscription.status = data["status"]
            db.session.commit()
        return subscription
    @app.route("/assinar/<plan>", methods=["GET", "POST"])
    def subscribe(plan):
        if plan not in PLANS:
            abort(404)
        if not os.environ.get("MP_ACCESS_TOKEN"):
            return render_template("subscribe.html", plan_key=plan, plan=PLANS[plan], unavailable=True), 503
        if request.method == "POST":
            email = request.form.get("email", "").strip().lower()
            if "@" not in email or len(email) > 254:
                flash("Informe um e-mail válido.", "error")
            elif db.session.scalar(select(User).where(User.email == email)):
                flash("Já existe uma conta com este e-mail. Use “Sou cliente” para entrar.", "error")
            else:
                subscription = Subscription(token=secrets.token_urlsafe(32), plan=plan, email=email)
                db.session.add(subscription)
                db.session.commit()
                # O Mercado Pago não aceita, em várias configurações, URLs locais
                # (localhost/127.0.0.1) como URL de retorno. Em desenvolvimento o
                # webhook continua sendo suficiente para sincronizar o pagamento;
                # em produção, PUBLIC_URL deve ser uma URL pública HTTPS.
                base = os.environ.get("PUBLIC_URL", "").strip().rstrip("/")
                payload = {
                    "reason": f"{app.config['APP_NAME']} — plano {PLANS[plan]['name']}",
                    "external_reference": subscription.token,
                    "payer_email": email,
                    "status": "pending",
                    "auto_recurring": {"frequency": 1, "frequency_type": "months", "transaction_amount": float(PLANS[plan]["price"]), "currency_id": "BRL"},
                }
                if base:
                    payload["back_url"] = base + url_for("subscription_return", ref=subscription.token)
                try:
                    created = mp_request("POST", "/preapproval", payload)
                    if not created.get("id") or not created.get("init_point"):
                        raise RuntimeError("Mercado Pago não retornou um checkout válido.")
                    subscription.mp_preapproval_id = created["id"]
                    db.session.commit()
                    return render_template("subscribe_redirect.html", checkout_url=created["init_point"])
                except Exception:
                    db.session.rollback()
                    app.logger.exception("Falha ao criar assinatura no Mercado Pago")
                    flash("Não foi possível iniciar o pagamento agora. Tente novamente em instantes.", "error")
        return render_template("subscribe.html", plan_key=plan, plan=PLANS[plan], unavailable=False)
    @app.get("/assinatura/retorno")
    def subscription_return():
        subscription = db.session.scalar(select(Subscription).where(Subscription.token == request.args.get("ref", ""))) or abort(404)
        if subscription.mp_preapproval_id:
            try:
                sync_subscription(subscription)
            except Exception:
                db.session.rollback()
                app.logger.exception("Falha ao consultar a assinatura no Mercado Pago")
        if subscription.status == "authorized":
            return redirect(url_for("register_subscriber", token=subscription.token) if not subscription.user_id else url_for("login"))
        return render_template("subscribe_status.html", subscription=subscription, plan=PLANS[subscription.plan])
    @app.post("/webhooks/mercadopago")
    def mercadopago_webhook():
        payload = request.get_json(silent=True) or {}
        data_id = str(request.args.get("data.id") or (payload.get("data") or {}).get("id") or "")
        secret = os.environ.get("MP_WEBHOOK_SECRET")
        if secret:
            parts = dict(item.split("=", 1) for item in request.headers.get("x-signature", "").split(",") if "=" in item)
            manifest = f"id:{data_id.lower()};request-id:{request.headers.get('x-request-id', '')};ts:{parts.get('ts', '')};"
            expected = hmac.new(secret.encode(), manifest.encode(), hashlib.sha256).hexdigest()
            if not secrets.compare_digest(expected, parts.get("v1", "")):
                abort(401)
        subscription = db.session.scalar(select(Subscription).where(Subscription.mp_preapproval_id == data_id)) if data_id else None
        if subscription:
            try:
                sync_subscription(subscription)
            except Exception:
                db.session.rollback()
                app.logger.exception("Falha ao processar webhook do Mercado Pago")
                return "", 500
        return "", 200
    @app.route("/cadastro/<token>", methods=["GET", "POST"])
    def register_subscriber(token):
        subscription = db.session.scalar(select(Subscription).where(Subscription.token == token))
        if not subscription or subscription.status != "authorized" or subscription.user_id:
            abort(410, "Este link de cadastro é inválido, já foi utilizado ou o pagamento ainda não foi confirmado.")
        if request.method == "POST":
            name = request.form.get("name", "").strip()
            username = request.form.get("username", "").strip().lower()
            password = request.form.get("password", "")
            try:
                if not name or len(name) > 120:
                    raise ValueError("Informe seu nome com até 120 caracteres.")
                if not username or len(username) > 50 or not username.replace("_", "").replace(".", "").isalnum():
                    raise ValueError("Use somente letras, números, ponto ou sublinhado no nome de usuário.")
                if len(password) < 10:
                    raise ValueError("A senha deve ter pelo menos 10 caracteres.")
                if db.session.scalar(select(User).where(or_(User.email == subscription.email, User.username == username))):
                    raise ValueError("O e-mail ou nome de usuário já está em uso.")
                user = User(name=name, username=username, email=subscription.email, password=generate_password_hash(password), plan=subscription.plan)
                db.session.add(user)
                db.session.flush()
                subscription.user_id = user.id
                db.session.add(UserAccessLog(user_id=user.id))
                db.session.commit()
                session.clear()
                session["user_id"] = user.id
                flash("Cadastro concluído. Boas-vindas!", "success")
                return redirect(url_for("index"))
            except ValueError as error:
                db.session.rollback()
                flash(str(error), "error")
        return render_template("register_subscriber.html", subscription=subscription, plan=PLANS[subscription.plan])
    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            identity = request.form.get("email", "").strip().lower()
            user = db.session.scalar(select(User).where(or_(User.email == identity, User.username == identity)))
            if user and check_password_hash(user.password, request.form.get("password", "")):
                subscription = db.session.scalar(select(Subscription).where(Subscription.user_id == user.id).order_by(Subscription.id.desc()))
                if subscription and subscription.status in ("cancelled", "paused"):
                    flash("Sua assinatura está inativa. Assine novamente para voltar a acessar.", "error")
                    return render_template("login.html")
                session.clear()
                session["user_id"] = user.id
                db.session.add(UserAccessLog(user_id=user.id))
                db.session.commit()
                return redirect(url_for("index"))
            flash("E-mail ou senha incorretos.", "error")
        return render_template("login.html")
    @app.get("/account")
    @login_required
    def account():
        user = db.session.get(User, session["user_id"])
        month_start = datetime(date.today().year, date.today().month, 1, tzinfo=timezone.utc)
        ai_used = db.session.scalar(select(func.count()).select_from(ClaudeAnalysis).where(ClaudeAnalysis.user_id == user.id, ClaudeAnalysis.created_at >= month_start))
        now = datetime.now(timezone.utc)
        links_active = sum(1 for link in db.session.scalars(select(SharedReport).where(SharedReport.user_id == user.id)) if utc_datetime(link.expires_at) > now)
        return render_template("account.html", user=user, ai_used=ai_used, links_active=links_active)
    @app.route("/invite/<token>", methods=["GET", "POST"])
    def accept_invitation(token):
        invitation = db.session.scalar(select(UserInvitation).where(UserInvitation.token == token))
        now = datetime.now(timezone.utc)
        if not invitation or invitation.used_at or utc_datetime(invitation.expires_at) <= now:
            abort(410, "Este convite expirou ou já foi utilizado.")
        if request.method == "POST":
            name = request.form.get("name", "").strip()
            username = request.form.get("username", "").strip().lower()
            email = request.form.get("email", "").strip().lower()
            password = request.form.get("password", "")
            try:
                if not name or len(name) > 120:
                    raise ValueError("Informe seu nome com até 120 caracteres.")
                if not username or len(username) > 50 or not username.replace("_", "").replace(".", "").isalnum():
                    raise ValueError("Use somente letras, números, ponto ou sublinhado no nome de usuário.")
                if "@" not in email or len(email) > 254:
                    raise ValueError("Informe um e-mail válido.")
                if len(password) < 10:
                    raise ValueError("A senha deve ter pelo menos 10 caracteres.")
                if db.session.scalar(select(User).where(or_(User.email == email, User.username == username))):
                    raise ValueError("O e-mail ou nome de usuário já está em uso.")
                user = User(name=name, username=username, email=email, password=generate_password_hash(password), is_admin=invitation.is_admin)
                invitation.used_at = now
                db.session.add(user)
                db.session.flush()
                db.session.add(UserAccessLog(user_id=user.id))
                db.session.commit()
                session.clear()
                session["user_id"] = user.id
                flash("Cadastro concluído. Boas-vindas!", "success")
                return redirect(url_for("index"))
            except ValueError as error:
                db.session.rollback()
                flash(str(error), "error")
        return render_template("accept_invitation.html", invitation=invitation)
    @app.route("/admin/users", methods=["GET", "POST"])
    @admin_required
    def admin_users():
        if request.method == "POST":
            name = request.form.get("name", "").strip()
            username = request.form.get("username", "").strip().lower()
            email = request.form.get("email", "").strip().lower()
            password = request.form.get("password", "")
            role = request.form.get("role", "user")
            try:
                if not name or len(name) > 120:
                    raise ValueError("Informe o nome com até 120 caracteres.")
                if not username or len(username) > 50 or not username.replace("_", "").replace(".", "").isalnum():
                    raise ValueError("Use somente letras, números, ponto ou sublinhado no nome de usuário.")
                if "@" not in email or len(email) > 254:
                    raise ValueError("Informe um e-mail válido.")
                if len(password) < 10:
                    raise ValueError("A senha deve ter pelo menos 10 caracteres.")
                if role not in ("user", "admin"):
                    raise ValueError("Nível de permissão inválido.")
                if db.session.scalar(select(User).where(or_(User.email == email, User.username == username))):
                    raise ValueError("O e-mail ou nome de usuário já está em uso.")
                db.session.add(User(name=name, username=username, email=email, password=generate_password_hash(password), is_admin=role == "admin"))
                db.session.commit()
                flash("Usuário criado.", "success")
            except ValueError as error:
                db.session.rollback()
                flash(str(error), "error")
            return redirect(url_for("admin_users"))
        users = db.session.scalars(select(User).order_by(User.name, User.username, User.email)).all()
        last_accesses = dict(db.session.execute(select(UserAccessLog.user_id, func.max(UserAccessLog.accessed_at)).group_by(UserAccessLog.user_id)).all())
        invitations = db.session.scalars(select(UserInvitation).where(UserInvitation.used_at.is_(None)).order_by(UserInvitation.expires_at.desc())).all()
        now = datetime.now(timezone.utc)
        invitation_items = [{"invitation": invitation, "expired": utc_datetime(invitation.expires_at) <= now} for invitation in invitations]
        return render_template("admin_users.html", users=users, invitations=invitation_items, last_accesses=last_accesses)
    @app.post("/admin/invitations")
    @admin_required
    def create_invitation():
        role = request.form.get("role", "user")
        try:
            expires_in = int(request.form.get("expires_in", ""))
            if role not in ("user", "admin"):
                raise ValueError("Nível de permissão inválido.")
            if expires_in not in (24, 48, 72, 168):
                raise ValueError("Selecione um prazo de validade válido.")
            invitation = UserInvitation(email="", is_admin=role == "admin", token=secrets.token_urlsafe(32), invited_by_id=g.current_user.id, expires_at=datetime.now(timezone.utc) + timedelta(hours=expires_in))
            db.session.add(invitation)
            db.session.commit()
            flash("Convite criado. Copie o link abaixo e envie ao usuário.", "success")
        except (TypeError, ValueError) as error:
            db.session.rollback()
            flash(str(error), "error")
        return redirect(url_for("admin_users"))
    @app.post("/admin/users/<int:id>/plan")
    @admin_required
    def update_plan(id):
        user = db.session.get(User, id) or abort(404)
        plan = request.form.get("plan", "")
        if plan not in PLANS:
            flash("Selecione um plano válido.", "error")
        else:
            user.plan = plan
            db.session.commit()
            flash(f"Plano de {user.name or user.username or user.email} alterado para {PLANS[plan]['name']}.", "success")
        return redirect(url_for("admin_users"))
    @app.post("/account/profile")
    @login_required
    def update_profile():
        user = db.session.get(User, session["user_id"])
        name = request.form.get("name", "").strip()
        username = request.form.get("username", "").strip().lower()
        email = request.form.get("email", "").strip().lower()
        avatar = request.files.get("avatar")
        if not name or len(name) > 120:
            flash("Informe seu nome com até 120 caracteres.", "error")
        elif not username or len(username) > 50 or not username.replace("_", "").replace(".", "").isalnum():
            flash("Use somente letras, números, ponto ou sublinhado no nome de usuário.", "error")
        elif "@" not in email or len(email) > 254:
            flash("Informe um e-mail válido.", "error")
        elif db.session.scalar(select(User).where(User.id != user.id, or_(User.email == email, User.username == username))):
            flash("O e-mail ou nome de usuário já está em uso.", "error")
        else:
            try:
                if avatar and avatar.filename:
                    stored, _mimetype = save_avatar(avatar, app.config["AVATAR_UPLOAD_FOLDER"])
                    old_avatar = user.avatar_file
                    user.avatar_file = stored
                    if old_avatar:
                        try:
                            os.remove(os.path.join(app.config["AVATAR_UPLOAD_FOLDER"], old_avatar))
                        except OSError:
                            pass
            except ValueError as error:
                flash(str(error), "error")
                return redirect(url_for("account"))
            user.name, user.username, user.email = name, username, email
            db.session.commit()
            flash("Informações atualizadas.", "success")
        return redirect(url_for("account"))
    @app.post("/account/avatar/remove")
    @login_required
    def remove_avatar():
        user = db.session.get(User, session["user_id"])
        if user.avatar_file:
            try:
                os.remove(os.path.join(app.config["AVATAR_UPLOAD_FOLDER"], user.avatar_file))
            except OSError:
                pass
            user.avatar_file = None
            db.session.commit()
        return redirect(url_for("account"))
    @app.get("/account/avatar")
    @login_required
    def avatar():
        user = db.session.get(User, session["user_id"])
        if not user.avatar_file:
            abort(404)
        extension = os.path.splitext(user.avatar_file)[1].lower()
        path = os.path.join(app.config["AVATAR_UPLOAD_FOLDER"], user.avatar_file)
        if not os.path.isfile(path):
            abort(404)
        return send_file(path, mimetype=AVATAR_TYPES.get(extension, "application/octet-stream"))
    @app.post("/account/password")
    @login_required
    def update_password():
        user = db.session.get(User, session["user_id"])
        current = request.form.get("current_password", "")
        new = request.form.get("new_password", "")
        if not check_password_hash(user.password, current):
            flash("A senha atual está incorreta.", "error")
        elif len(new) < 10:
            flash("A nova senha deve ter pelo menos 10 caracteres.", "error")
        elif new != request.form.get("password_confirmation", ""):
            flash("A confirmação da nova senha não corresponde.", "error")
        else:
            user.password = generate_password_hash(new)
            db.session.commit()
            flash("Senha alterada com sucesso.", "success")
        return redirect(url_for("account"))
    @app.post("/logout")
    def logout():
        session.clear()
        return redirect(url_for("login"))
    @app.get("/")
    def index():
        if not session.get("user_id"):
            return render_template("landing.html")
        try:
            selected = datetime.strptime(request.args.get("month", date.today().strftime("%Y-%m")), "%Y-%m").date()
        except ValueError:
            abort(400)
        entries = db.session.scalars(select(Entry).where(Entry.user_id == session["user_id"], Entry.month == selected, Entry.deleted_at.is_(None))).all()
        entries.sort(key=lambda entry: entry.description.casefold())
        group_definitions = (
            ("receitas", "Receitas", "Valores previstos para entrar", lambda entry: entry.kind == "receita"),
            ("fixas", "Despesas fixas", "Contas recorrentes e compromissos", lambda entry: entry.kind == "despesa" and entry.category != "dia_a_dia"),
            ("dia-a-dia", "Despesas do dia a dia", "Gastos pontuais realizados no mês", lambda entry: entry.kind == "despesa" and entry.category == "dia_a_dia"),
        )
        entry_groups = []
        for key, title, subtitle, predicate in group_definitions:
            items = [entry for entry in entries if predicate(entry)]
            if key == "dia-a-dia":
                items.sort(key=lambda entry: (entry.expense_date is None, entry.expense_date, entry.description.casefold()))
            entry_groups.append({"key": key, "title": title, "subtitle": subtitle, "entries": items, "total": sum((entry.amount for entry in items), Decimal(0))})
        expenses = sum((e.amount for e in entries if e.kind == "despesa"), Decimal(0))
        income = sum((e.amount for e in entries if e.kind == "receita"), Decimal(0))
        paid = sum((e.amount for e in entries if e.kind == "despesa" and e.paid), Decimal(0))
        received = sum((e.amount for e in entries if e.kind == "receita" and e.paid), Decimal(0))
        analyses = db.session.scalars(select(ClaudeAnalysis).where(ClaudeAnalysis.user_id == session["user_id"]).order_by(ClaudeAnalysis.created_at.desc()).limit(30)).all()
        return render_template("index.html", entries=entries, entry_groups=entry_groups, selected=selected, expenses=expenses, income=income, paid=paid, received=received, analyses=analyses)
    @app.get("/reports")
    @login_required
    def reports():
        try:
            selected = datetime.strptime(request.args.get("month", date.today().strftime("%Y-%m")), "%Y-%m").date()
        except ValueError:
            abort(400)
        year_entries = db.session.scalars(select(Entry).where(Entry.user_id == session["user_id"], Entry.deleted_at.is_(None), Entry.month >= date(selected.year, 1, 1), Entry.month <= date(selected.year, 12, 1))).all()
        entries = sorted((entry for entry in year_entries if entry.month == selected), key=lambda entry: entry.description.casefold())
        definitions = (
            ("receitas", "Receitas", "Valores previstos para entrar", lambda entry: entry.kind == "receita"),
            ("fixas", "Despesas fixas", "Contas recorrentes e compromissos", lambda entry: entry.kind == "despesa" and entry.category != "dia_a_dia"),
            ("dia-a-dia", "Despesas do dia a dia", "Gastos pontuais realizados no mês", lambda entry: entry.kind == "despesa" and entry.category == "dia_a_dia"),
        )
        entry_groups = []
        for key, title, subtitle, predicate in definitions:
            items = [entry for entry in entries if predicate(entry)]
            if key == "dia-a-dia":
                items.sort(key=lambda entry: (entry.expense_date is None, entry.expense_date, entry.description.casefold()))
            entry_groups.append({"key": key, "title": title, "subtitle": subtitle, "entries": items, "total": sum((entry.amount for entry in items), Decimal(0))})
        chart = []
        for month_number in range(1, 13):
            monthly = [entry for entry in year_entries if entry.month.month == month_number]
            chart.append({"month": MONTHS[month_number-1][:3], "income": sum((entry.amount for entry in monthly if entry.kind == "receita"), Decimal(0)), "expenses": sum((entry.amount for entry in monthly if entry.kind == "despesa"), Decimal(0))})
        chart_max = max((max(point["income"], point["expenses"]) for point in chart), default=Decimal(1)) or Decimal(1)
        share_token = request.args.get("share")
        share_url = url_for("shared_report", token=share_token, _external=True) if share_token else None
        return render_template("reports.html", selected=selected, entry_groups=entry_groups, chart=chart, chart_max=chart_max, share_url=share_url)
    def report_data(month_value):
        try:
            selected = datetime.strptime(month_value, "%Y-%m").date()
        except (TypeError, ValueError):
            abort(400, "Mês de referência inválido.")
        entries = db.session.scalars(select(Entry).where(Entry.user_id == session["user_id"], Entry.month == selected, Entry.deleted_at.is_(None))).all()
        entries.sort(key=lambda entry: entry.description.casefold())
        income = sum((entry.amount for entry in entries if entry.kind == "receita"), Decimal(0))
        expenses = sum((entry.amount for entry in entries if entry.kind == "despesa"), Decimal(0))
        return selected, entries, income, expenses
    @app.get("/reports/export/excel")
    @login_required
    def export_month_excel():
        selected, entries, income, expenses = report_data(request.args.get("month"))
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Relatório mensal"
        sheet.merge_cells("A1:F1")
        sheet["A1"] = f"Relatório financeiro — {MONTHS[selected.month - 1]} de {selected.year}"
        sheet["A1"].font = Font(bold=True, size=15, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="176B5D")
        sheet["A1"].alignment = Alignment(horizontal="center")
        for label, value in (("Receitas", income), ("Despesas", expenses), ("Saldo previsto", income - expenses)):
            sheet.append([label, value])
            sheet.cell(sheet.max_row, 1).font = Font(bold=True)
            sheet.cell(sheet.max_row, 2).number_format = 'R$ #,##0.00'
        sheet.append([])
        headers = ("Descrição", "Tipo", "Tipo de despesa", "Data", "Valor", "Situação")
        sheet.append(headers)
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="102A43")
        for entry in entries:
            row = entry_report_row(entry)
            sheet.append((*row[:4], float(row[4]), row[5]))
            sheet.cell(sheet.max_row, 5).number_format = 'R$ #,##0.00'
        sheet.freeze_panes = f"A{header_row + 1}"
        widths = (32, 14, 20, 15, 16, 14)
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[chr(64 + column)].width = width
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        filename = f"relatorio-financeiro-{selected.strftime('%Y-%m')}.xlsx"
        return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=filename)
    @app.get("/reports/export/pdf")
    @login_required
    def export_month_pdf():
        selected, entries, income, expenses = report_data(request.args.get("month"))
        output = BytesIO()
        document = SimpleDocTemplate(output, pagesize=A4, rightMargin=1.3 * cm, leftMargin=1.3 * cm, topMargin=1.3 * cm, bottomMargin=1.3 * cm)
        styles = getSampleStyleSheet()
        story = [Paragraph(app.config["APP_NAME"], styles["Title"]), Paragraph("Relatório financeiro", styles["Heading2"]), Paragraph(f"{MONTHS[selected.month - 1]} de {selected.year}", styles["Heading2"]), Spacer(1, 0.35 * cm)]
        summary = Table([["Receitas", format_brl(income)], ["Despesas", format_brl(expenses)], ["Saldo previsto", format_brl(income - expenses)]], colWidths=(9 * cm, 7 * cm))
        summary.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#E6F3F0")), ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#176B5D")), ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9DDDA")), ("PADDING", (0, 0), (-1, -1), 7)]))
        story.extend([summary, Spacer(1, 0.5 * cm), Paragraph("Lançamentos", styles["Heading2"])])
        rows = [["Descrição", "Tipo", "Categoria", "Data", "Valor", "Situação"]]
        for entry in entries:
            row = entry_report_row(entry)
            rows.append([Paragraph(escape(str(value)), styles["BodyText"]) for value in (*row[:4], format_brl(row[4]), row[5])])
        table = Table(rows, colWidths=(5.0 * cm, 2.1 * cm, 2.4 * cm, 2.1 * cm, 2.4 * cm, 2.4 * cm), repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#102A43")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D8DFE3")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFA")]), ("PADDING", (0, 0), (-1, -1), 5)]))
        story.append(table)
        document.build(story)
        output.seek(0)
        filename = f"relatorio-financeiro-{selected.strftime('%Y-%m')}.pdf"
        return send_file(output, mimetype="application/pdf", as_attachment=True, download_name=filename)
    @app.post("/reports/share")
    @login_required
    def create_shared_report():
        selected, _entries, _income, _expenses = report_data(request.form.get("month"))
        share_limit = g.current_user.plan_info["share_links"]
        if share_limit is not None:
            now = datetime.now(timezone.utc)
            active = sum(1 for link in db.session.scalars(select(SharedReport).where(SharedReport.user_id == session["user_id"])) if utc_datetime(link.expires_at) > now)
            if active >= share_limit:
                flash("O plano Básico não inclui links de compartilhamento." if share_limit == 0 else f"Você atingiu o limite de {share_limit} links ativos do plano {g.current_user.plan_info['name']}. Aguarde um link expirar ou mude de plano.", "error")
                return redirect(url_for("reports", month=selected.strftime("%Y-%m")))
        payload = json.dumps({"user_id": session["user_id"], "month": selected.strftime("%Y-%m")}).encode()
        token = app.config["REPORT_SHARE_CIPHER"].encrypt(payload).decode()
        created_at = datetime.now(timezone.utc)
        db.session.add(SharedReport(user_id=session["user_id"], month=selected, token=token, created_at=created_at, expires_at=created_at + timedelta(hours=3)))
        db.session.commit()
        return redirect(url_for("reports", month=selected.strftime("%Y-%m"), share=token))
    @app.get("/shared-links")
    @login_required
    def shared_links():
        now = datetime.now(timezone.utc)
        links = db.session.scalars(select(SharedReport).where(SharedReport.user_id == session["user_id"]).order_by(SharedReport.created_at.desc())).all()
        items = []
        for link in links:
            expires_at = link.expires_at.replace(tzinfo=timezone.utc) if link.expires_at.tzinfo is None else link.expires_at
            seconds = max(0, int((expires_at - now).total_seconds()))
            hours, remainder = divmod(seconds, 3600)
            minutes = remainder // 60
            items.append({"link": link, "url": url_for("shared_report", token=link.token, _external=True), "active": seconds > 0, "remaining": f"{hours}h {minutes:02d}min" if seconds else "Expirado"})
        return render_template("shared_links.html", links=items)
    def shared_report_context(token):
        try:
            payload = json.loads(app.config["REPORT_SHARE_CIPHER"].decrypt(token.encode(), ttl=3 * 60 * 60))
            user_id = int(payload["user_id"])
            selected = datetime.strptime(payload["month"], "%Y-%m").date()
        except (InvalidToken, KeyError, TypeError, ValueError, json.JSONDecodeError):
            abort(410, "Este link de relatório expirou ou é inválido.")
        sender = db.session.get(User, user_id)
        if not sender:
            abort(410, "Este link de relatório expirou ou é inválido.")
        return user_id, selected, sender
    @app.get("/shared-report/<token>")
    def shared_report(token):
        user_id, selected, sender = shared_report_context(token)
        entries = db.session.scalars(select(Entry).where(Entry.user_id == user_id, Entry.month == selected, Entry.deleted_at.is_(None))).all()
        entries.sort(key=lambda entry: entry.description.casefold())
        groups = (
            ("Receitas", lambda entry: entry.kind == "receita"),
            ("Despesas fixas", lambda entry: entry.kind == "despesa" and entry.category != "dia_a_dia"),
            ("Despesas do dia a dia", lambda entry: entry.kind == "despesa" and entry.category == "dia_a_dia"),
        )
        entry_groups = []
        for title, predicate in groups:
            items = [entry for entry in entries if predicate(entry)]
            if title == "Despesas do dia a dia":
                items.sort(key=lambda entry: (entry.expense_date is None, entry.expense_date, entry.description.casefold()))
            entry_groups.append({"title": title, "entries": items})
        income = sum((entry.amount for entry in entries if entry.kind == "receita"), Decimal(0))
        expenses = sum((entry.amount for entry in entries if entry.kind == "despesa"), Decimal(0))
        received = sum((entry.amount for entry in entries if entry.kind == "receita" and entry.paid), Decimal(0))
        paid = sum((entry.amount for entry in entries if entry.kind == "despesa" and entry.paid), Decimal(0))
        sender_name = sender.name or sender.username or sender.email
        return render_template("shared_report.html", selected=selected, entry_groups=entry_groups, income=income, expenses=expenses, received=received, paid=paid, sender_name=sender_name, share_token=token, shared_report=True)
    @app.get("/shared-report/<token>/receipt/<int:id>")
    def shared_receipt(token, id):
        user_id, selected, _sender = shared_report_context(token)
        entry = db.session.scalar(select(Entry).where(Entry.id == id, Entry.user_id == user_id, Entry.month == selected, Entry.deleted_at.is_(None)))
        if not entry or not entry.paid or not entry.receipt_file:
            abort(404)
        path = os.path.join(app.config["RECEIPT_UPLOAD_FOLDER"], entry.receipt_file)
        if not os.path.isfile(path):
            abort(404)
        return send_file(path, mimetype=entry.receipt_mimetype, download_name=entry.receipt_name, as_attachment=request.args.get("download") == "1")
    @app.get("/planejamento-anual")
    @login_required
    def annual_planning():
        try:
            selected = datetime.strptime(request.args.get("month", date.today().strftime("%Y-%m")), "%Y-%m").date()
        except ValueError:
            abort(400)
        annual = db.session.scalars(select(Entry).where(Entry.user_id == session["user_id"], Entry.deleted_at.is_(None), Entry.month >= date(selected.year, 1, 1), Entry.month <= date(selected.year, 12, 1))).all()
        grid = {}
        for entry in annual:
            grid.setdefault((entry.description, entry.kind), {}).setdefault(entry.month.month, []).append(entry)
        annual_rows = sorted(grid.items(), key=lambda item: item[0][0].casefold())
        return render_template("annual_planning.html", selected=selected, annual_rows=annual_rows)
    def owned(id):
        return db.session.scalar(select(Entry).where(Entry.id == id, Entry.user_id == session["user_id"], Entry.deleted_at.is_(None))) or abort(404)
    @app.route("/entry/new", methods=["GET", "POST"])
    @app.route("/entry/<int:id>/edit", methods=["GET", "POST"])
    @login_required
    def edit(id=None):
        entry = owned(id) if id else None
        if request.method == "GET":
            target_month = entry.month if entry else date.today().replace(day=1)
            return redirect(url_for("index", month=target_month.strftime("%Y-%m"), modal="entry", edit=id or ""))
        if request.method == "POST":
            try:
                description = request.form.get("description", "").strip()
                if not description or len(description) > 150:
                    raise ValueError("Informe uma descrição de até 150 caracteres.")
                kind = request.form.get("kind")
                if kind not in ("receita", "despesa"):
                    raise ValueError("Tipo inválido.")
                category = request.form.get("category", "fixa") if kind == "despesa" else "receita"
                if category not in ("fixa", "dia_a_dia", "receita"):
                    raise ValueError("Categoria inválida.")
                month = datetime.strptime(request.form.get("month", ""), "%Y-%m").date()
                entry_date = parse_date(request.form.get("entry_date") or request.form.get("due"))
                if category == "dia_a_dia":
                    if not entry_date:
                        raise ValueError("Informe a data do gasto.")
                    month = entry_date.replace(day=1)
                amount = money(request.form.get("amount"))
                repeat_value = request.form.get("repeat", "1")
                repeat = 1 if entry or category != "fixa" or repeat_value == "never" else int(repeat_value)
                if not 1 <= repeat <= 12:
                    raise ValueError("Repetição deve ser de 1 a 12 meses.")
                receipt = request.files.get("receipt")
                receipt_data = None
                if receipt and receipt.filename:
                    if kind != "despesa" or category != "fixa":
                        raise ValueError("Comprovantes estão disponíveis somente para despesas fixas.")
                    receipt_data = save_receipt(receipt, app.config["RECEIPT_UPLOAD_FOLDER"])
                old_receipt = entry.receipt_file if entry and (receipt_data or category != "fixa") else None
                if entry and category != "fixa":
                    entry.receipt_file = entry.receipt_name = entry.receipt_mimetype = None
                    entry.receipt_uploaded_at = None
                for offset in range(repeat):
                    month_index = month.year*12 + month.month-1+offset
                    target = date(month_index//12, month_index%12+1, 1)
                    obj = entry or Entry(user_id=session["user_id"])
                    obj.description, obj.kind, obj.category, obj.month, obj.amount = description, kind, category, target, amount
                    obj.expense_date = entry_date if category == "dia_a_dia" else None
                    obj.due = None if category == "dia_a_dia" else entry_date if not offset or not entry_date else date(target.year, target.month, min(entry_date.day, calendar.monthrange(target.year, target.month)[1]))
                    obj.paid = request.form.get("paid") == "on" if offset == 0 else False
                    if receipt_data and offset == 0:
                        obj.receipt_file, obj.receipt_name, obj.receipt_mimetype = receipt_data
                        obj.receipt_uploaded_at = datetime.now(timezone.utc)
                    db.session.add(obj)
                db.session.commit()
                if old_receipt:
                    try:
                        os.remove(os.path.join(app.config["RECEIPT_UPLOAD_FOLDER"], old_receipt))
                    except FileNotFoundError:
                        pass
                flash("Lançamento salvo.", "success")
                return redirect(url_for("index", month=month.strftime("%Y-%m")))
            except (ValueError, TypeError) as exc:
                db.session.rollback()
                flash(str(exc), "error")
        target_month = entry.month if entry else date.today().replace(day=1)
        return redirect(url_for("index", month=target_month.strftime("%Y-%m"), modal="entry", edit=id or ""))
    @app.get("/entry/<int:id>/receipt")
    @login_required
    def receipt(id):
        entry = owned(id)
        if entry.category != "fixa" or not entry.receipt_file:
            abort(404)
        path = os.path.join(app.config["RECEIPT_UPLOAD_FOLDER"], entry.receipt_file)
        if not os.path.isfile(path):
            abort(404)
        return send_file(path, mimetype=entry.receipt_mimetype, download_name=entry.receipt_name, as_attachment=request.args.get("download") == "1")
    @app.post("/entry/<int:id>/toggle")
    @login_required
    def toggle(id):
        e = owned(id)
        e.paid = not e.paid
        db.session.commit()
        return redirect(url_for("index", month=e.month.strftime("%Y-%m")))
    @app.post("/entry/<int:id>/delete")
    @login_required
    def delete(id):
        e = owned(id)
        month = e.month.strftime("%Y-%m")
        e.deleted_at = datetime.now(timezone.utc)
        db.session.commit()
        flash("Lançamento excluído.", "success")
        return redirect(url_for("index", month=month))
    @app.post("/api/suggestions")
    @login_required
    def financial_suggestions():
        api_key = os.environ.get("ANTHROPIC_API_KEY")
        if not api_key:
            return jsonify(error="O serviço de sugestões financeiras ainda não foi configurado."), 503
        try:
            selected = datetime.strptime(request.form.get("month", ""), "%Y-%m").date()
        except ValueError:
            return jsonify(error="Mês de referência inválido."), 400
        current_user = g.current_user
        usage_start = datetime(date.today().year, date.today().month, 1, tzinfo=timezone.utc)
        usage = db.session.scalar(select(db.func.count()).select_from(ClaudeAnalysis).where(ClaudeAnalysis.user_id == current_user.id, ClaudeAnalysis.created_at >= usage_start))
        ai_limit = current_user.plan_info["ai"]
        if usage >= ai_limit:
            return jsonify(error=f"Você atingiu o limite de {ai_limit} sugestões de IA do plano {current_user.plan_info['name']} neste mês."), 429
        entries = db.session.scalars(select(Entry).where(Entry.user_id == session["user_id"], Entry.month == selected, Entry.deleted_at.is_(None))).all()
        pending = [entry for entry in entries if not entry.paid]
        income = sum((entry.amount for entry in entries if entry.kind == "receita"), Decimal(0))
        expenses = sum((entry.amount for entry in entries if entry.kind == "despesa"), Decimal(0))
        received = sum((entry.amount for entry in entries if entry.kind == "receita" and entry.paid), Decimal(0))
        paid = sum((entry.amount for entry in entries if entry.kind == "despesa" and entry.paid), Decimal(0))
        today = date.today()
        month_end = date(selected.year, selected.month, calendar.monthrange(selected.year, selected.month)[1])
        if selected.year == today.year and selected.month == today.month:
            remaining_days = (month_end - today).days + 1
        elif selected > today.replace(day=1):
            remaining_days = month_end.day
        else:
            remaining_days = 0
        pending_data = [claude_entry_data(entry) for entry in pending]
        context = {
            "periodo_referencia": selected.strftime("%Y-%m"),
            "resumo": {
                "receitas_previstas": str(income),
                "despesas_previstas": str(expenses),
                "saldo_previsto": str(income - expenses),
                "receitas_recebidas": str(received),
                "despesas_pagas": str(paid),
                "saldo_atual": str(received - paid),
                "dias_restantes_no_periodo": remaining_days,
            },
            "lancamentos_pendentes": pending_data,
        }
        prompt = f"""Analise o orçamento usando exclusivamente o JSON abaixo.

Contrato dos lançamentos:
- tipo_lancamento é sempre "receita" ou "despesa".
- tipo_despesa só existe para tipo_lancamento="despesa": "fixa" para contas recorrentes/compromissos e "dia_a_dia" para gastos pontuais. Para receitas ele é null.
- tipo_data informa se a data é de recebimento, vencimento ou gasto.

Dados financeiros:
{json.dumps(context, ensure_ascii=False, indent=2)}

Responda em português do Brasil, de forma curta e prática. Em toda resposta, comece com “Previsão diária de gasto: R$ X por dia”, calculada a partir de saldo_atual dividido por dias_restantes_no_periodo. Não considere receitas apenas previstas como saldo disponível. Se o saldo atual for zero ou negativo, informe R$ 0,00 por dia e alerte para não assumir novos gastos; se não houver dias restantes, diga que não é possível fazer a previsão diária para um período encerrado. Em seguida, indique o próximo pagamento recomendado, depois uma ordem numerada para os demais itens e finalize com até três sugestões financeiras. Priorize contas vencidas, vencimentos próximos e despesas essenciais, mas deixe claro quando a descrição não permite determinar se algo é essencial. Não invente juros, renda ou datas. Não use Markdown além de listas simples. Inclua ao final: “Orientação geral; confirme multas, juros e prioridades antes de decidir.”"""
        try:
            from anthropic import Anthropic
            client = Anthropic(api_key=api_key, timeout=20.0, max_retries=1)
            model = os.environ.get("ANTHROPIC_MODEL", "claude-haiku-4-5-20251001")
            message = client.messages.create(model=model, max_tokens=500, system="Você é um assistente de organização financeira cauteloso. Use somente os dados fornecidos e nunca prometa resultados.", messages=[{"role": "user", "content": prompt}])
            suggestion = "\n".join(block.text for block in message.content if block.type == "text").strip()
            if not suggestion:
                raise ValueError("Resposta vazia")
            analysis = ClaudeAnalysis(user_id=session["user_id"], month=selected, content=suggestion, model=model)
            db.session.add(analysis)
            db.session.commit()
            return jsonify(suggestion=suggestion, saved=True, created_at=br_datetime(analysis.created_at))
        except Exception:
            db.session.rollback()
            app.logger.exception("Falha ao consultar a API do Claude")
            return jsonify(error="Não foi possível gerar as sugestões agora. Tente novamente em alguns instantes."), 502
    @app.route("/import", methods=["GET", "POST"])
    @login_required
    def import_excel():
        if request.method == "GET":
            return redirect(url_for("index", modal="import"))
        if request.method == "POST":
            file = request.files.get("file")
            try:
                if not file or not file.filename.lower().endswith(".xlsx"):
                    raise ValueError("Selecione um arquivo .xlsx.")
                year = int(request.form.get("year", date.today().year))
                if not 1900 <= year <= 2100:
                    raise ValueError("Ano deve estar entre 1900 e 2100.")
                parsed = read_excel(file, year)
                existing = db.session.scalars(select(Entry).where(Entry.user_id == session["user_id"], Entry.deleted_at.is_(None))).all()
                signature = lambda e: (e.description, e.kind, e.month, e.due, e.amount, e.paid)
                seen = {signature(e) for e in existing}
                count = 0
                for data in parsed:
                    e = Entry(user_id=session["user_id"], **data)
                    if signature(e) not in seen:
                        db.session.add(e)
                        seen.add(signature(e))
                        count += 1
                db.session.commit()
                flash(f"{count} lançamentos importados; {len(parsed)-count} duplicados ignorados.", "success")
                return redirect(url_for("index", month=parsed[0]["month"].strftime("%Y-%m")))
            except Exception as exc:
                db.session.rollback()
                flash(
                    import_error_message(exc) if isinstance(exc, ValueError) else "Não foi possível importar. Confira o arquivo e o formato do modelo ou baixe um novo modelo.",
                    "error import-error",
                )
        return redirect(url_for("index", modal="import"))
    @app.get("/model.xlsx")
    @login_required
    def model():
        wb = Workbook()
        ws = wb.active
        ws.title = "Finanças"
        ws.append(["Descrição", "Tipo", "Mês", "Vencimento", "Valor", "Pago"])
        ws.append(["Aluguel", "despesa", "2026-09", "10/09/2026", 800, "Não"])
        ws.append(["Salário", "receita", "2026-09", "05/09/2026", 2500, "Sim"])
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name="modelo-financas.xlsx")
    @app.cli.command("init-db")
    def init_db():
        db.create_all()
        click.echo("Banco de dados inicializado.")
    @app.cli.command("upgrade-db")
    def upgrade_db():
        db.create_all()
        columns = {column["name"] for column in inspect(db.engine).get_columns("entry")}
        with db.engine.begin() as connection:
            if "category" not in columns:
                connection.execute(text("ALTER TABLE entry ADD COLUMN category VARCHAR(20) NOT NULL DEFAULT 'fixa'"))
            if "expense_date" not in columns:
                connection.execute(text("ALTER TABLE entry ADD COLUMN expense_date DATE"))
            if "receipt_file" not in columns:
                connection.execute(text("ALTER TABLE entry ADD COLUMN receipt_file VARCHAR(80)"))
            if "receipt_name" not in columns:
                connection.execute(text("ALTER TABLE entry ADD COLUMN receipt_name VARCHAR(255)"))
            if "receipt_mimetype" not in columns:
                connection.execute(text("ALTER TABLE entry ADD COLUMN receipt_mimetype VARCHAR(100)"))
            if "receipt_uploaded_at" not in columns:
                connection.execute(text("ALTER TABLE entry ADD COLUMN receipt_uploaded_at DATETIME"))
            if "deleted_at" not in columns:
                connection.execute(text("ALTER TABLE entry ADD COLUMN deleted_at DATETIME"))
        user_columns = {column["name"] for column in inspect(db.engine).get_columns("user")}
        with db.engine.begin() as connection:
            if "name" not in user_columns:
                connection.execute(text("ALTER TABLE user ADD COLUMN name VARCHAR(120)"))
            if "username" not in user_columns:
                connection.execute(text("ALTER TABLE user ADD COLUMN username VARCHAR(50)"))
            if "avatar_file" not in user_columns:
                connection.execute(text("ALTER TABLE user ADD COLUMN avatar_file VARCHAR(80)"))
            if "is_admin" not in user_columns:
                connection.execute(text("ALTER TABLE user ADD COLUMN is_admin BOOLEAN NOT NULL DEFAULT 0"))
            if "ai_monthly_limit" not in user_columns:
                connection.execute(text("ALTER TABLE user ADD COLUMN ai_monthly_limit INTEGER NOT NULL DEFAULT 10"))
            if "plan" not in user_columns:
                connection.execute(text("ALTER TABLE user ADD COLUMN plan VARCHAR(20) NOT NULL DEFAULT 'basico'"))
                # Quem já usava o sistema mantém o acesso atual: administradores no Premium e os demais no Padrão.
                connection.execute(text("UPDATE user SET plan = 'premium' WHERE is_admin = 1"))
                connection.execute(text("UPDATE user SET plan = 'padrao' WHERE is_admin = 0"))
            connection.execute(text("UPDATE user SET is_admin = 1 WHERE username = 'chiapettaiago'"))
        click.echo("Estrutura do banco atualizada.")
    @app.cli.command("create-user")
    @click.argument("email")
    @click.password_option(confirmation_prompt=True)
    def create_user(email, password):
        if len(password) < 10:
            raise click.ClickException("Use uma senha com pelo menos 10 caracteres.")
        email = email.strip().lower()
        if "@" not in email or len(email) > 254:
            raise click.ClickException("E-mail inválido.")
        if db.session.scalar(select(User).where(User.email == email)):
            raise click.ClickException("Usuário já cadastrado.")
        db.session.add(User(email=email, password=generate_password_hash(password)))
        db.session.commit()
        click.echo("Usuário criado.")
    return app

app = create_app()
